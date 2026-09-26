"""
Turn the blinded audit CSV (written by e2_score.py) into an Excel workbook that is easy to rate.

    python build_audit_workbook.py results/audit_sheet.csv results/audit_rater1.xlsx

Sheets: "How to rate" (rules + an example), "Rating" (the items, with Y/N drop-downs),
"Progress" (live counts and consistency checks). No method names appear anywhere.
"""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
a = pd.read_csv(src)
C1, C2 = "is_it_in_the_photo(Y/N)", "would_a_patient_here_plausibly_want_or_mention_it(Y/N)"
scene_no = {iid: i + 1 for i, iid in enumerate(dict.fromkeys(a.image_id))}
a["scene"] = a.image_id.map(scene_no)
a = a.sort_values(["scene", "item"], kind="stable").reset_index(drop=True)

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F3864")
BAND = PatternFill("solid", fgColor="EEF3FA")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
wb = Workbook()

# ---------------- How to rate ----------------
h = wb.active
h.title = "How to rate"
lines = [
    ("E2 blinded audit: how to rate", "title"),
    ("", None),
    ("What you are doing", "h"),
    ("Each row is one object name that some method proposed for a photo. The method is hidden. "
     "You judge each name by looking at the photo. There are 50 photos and about 21 names per photo (1,045 rows).", None),
    ("", None),
    ("Steps", "h"),
    ("1. Go to the 'Rating' sheet. Rows are grouped by photo (Scene 1 to 50).", None),
    ("2. Click 'Open photo' in the first row of a scene. The photo opens in your browser. Keep it open beside Excel.", None),
    ("3. For every name in that scene, fill column D (yellow): Y or N from the drop-down.", None),
    ("4. Only if D is Y, fill column E: Y or N. If D is N, leave E empty (it turns grey).", None),
    ("5. Use column F for a short note only if you were unsure (optional).", None),
    ("6. Move to the next scene. Save often (Ctrl+S). The 'Progress' sheet shows how far you are and flags mistakes.", None),
    ("7. When Progress says 100% and 0 problems, save and close. Send the file back (or leave it in the folder).", None),
    ("", None),
    ("Column D: 'Can you point to this thing in the photo?'", "h"),
    ("Y = you can point to it, even if it is small, partly hidden or at the edge.", None),
    ("N = it is not there, OR you cannot tell, OR the name is not a thing you could point to "
     "(e.g. 'object', 'area', 'room', 'kitchen', 'furniture', 'group of people').", None),
    ("Accept close names: 'couch' for a sofa, 'tv' for a television, 'cup' for a mug. Plural is fine: 'sinks' = Y if one sink is there.", None),
    ("A picture of the thing (e.g. a poster showing a dog) does not count: N. A person counts as 'person', 'man', 'woman': Y if visible.", None),
    ("", None),
    ("Column E (only when D = Y): 'Would a person in this room who cannot move or speak plausibly want, use or talk about it?'", "h"),
    ("Imagine someone with severe paralysis lying or sitting in this room, choosing items on a screen to say what they need.", None),
    ("Y = they could ask for it, ask to use it, or mention it (water bottle, blanket, tv, window, toilet, phone, book, lamp, person).", None),
    ("N = it is a background part they would hardly ever name (floor, wall, ceiling, tiles, cabinet handle, electric outlet, shower head).", None),
    ("When in doubt, answer Y. Do not think about which method might have produced the name.", None),
    ("", None),
    ("Example (a bathroom photo with a sink, a mirror and a towel; no bathtub)", "h"),
]
r = 1
for text, kind in lines:
    c = h.cell(row=r, column=1, value=text)
    c.font = Font(name=F, size=14 if kind == "title" else 11, bold=kind in ("title", "h"))
    c.alignment = Alignment(wrap_text=True, vertical="top")
    h.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    h.row_dimensions[r].height = 22 if kind == "title" else 15 * max(1, -(-len(text) // 105))
    r += 1
ex = [("Item", "D: in photo?", "E: would they want it?", "Why"),
      ("towel", "Y", "Y", "Visible, and a person may ask for it"),
      ("mirror", "Y", "Y", "Visible, and they might mention it"),
      ("tiles", "Y", "N", "Visible, but just the wall surface"),
      ("bathtub", "N", "", "Not in this photo, so E stays empty"),
      ("object", "N", "", "Not a thing you can point to")]
for i, row in enumerate(ex):
    for j, v in enumerate(row):
        c = h.cell(row=r + i, column=1 + j, value=v)
        c.font = Font(name=F, bold=(i == 0), color="FFFFFF" if i == 0 else "000000")
        if i == 0:
            c.fill = HF
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
r += len(ex) + 1
h.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
h.cell(row=r, column=1, value="Do not open the folder 'audit_key_private' until both raters have finished: it reveals the methods.").font = Font(name=F, bold=True, color="C00000")
for col, w in zip("ABCD", (18, 16, 24, 62)):
    h.column_dimensions[col].width = w

# ---------------- Rating ----------------
s = wb.create_sheet("Rating")
hdr = ["Scene", "Photo", "Item (object name)", "D: In the photo? (Y/N)",
       "E: Would they want / mention it? (Y/N, only if D = Y)", "Notes (optional)", "item_uid", "image_id"]
for j, v in enumerate(hdr, 1):
    c = s.cell(row=1, column=j, value=v)
    c.font, c.fill = H, HF
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
s.row_dimensions[1].height = 45
n = len(a)
prev = None
for i, row in a.iterrows():
    rr = i + 2
    first = row.scene != prev
    prev = row.scene
    s.cell(row=rr, column=1, value=int(row.scene))
    if first:
        c = s.cell(row=rr, column=2, value="Open photo")
        c.hyperlink = row.url
        c.font = Font(name=F, color="0563C1", underline="single", bold=True)
    s.cell(row=rr, column=3, value=row["item"])
    for col, key in ((4, C1), (5, C2)):
        v = row.get(key)
        s.cell(row=rr, column=col, value=None if pd.isna(v) else str(v).strip().upper())
    s.cell(row=rr, column=7, value=row.item_uid)
    s.cell(row=rr, column=8, value=int(row.image_id))
    band = BAND if row.scene % 2 == 0 else None
    for col in range(1, 9):
        c = s.cell(row=rr, column=col)
        if col not in (2,):
            c.font = Font(name=F, color="808080" if col >= 7 else "000000")
        if col in (4, 5):
            c.fill = INPUT
            c.alignment = Alignment(horizontal="center")
        elif band is not None:
            c.fill = band
        if first:
            c.border = Border(top=Side(style="medium", color="1F3864"))
last = n + 1
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                    errorTitle="Y or N", error="Please choose Y or N from the list.")
s.add_data_validation(dv)
dv.add(f"D2:E{last}")
red = PatternFill("solid", fgColor="F4B6B6")        # rules added in priority order: red wins
s.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['AND($D2="N",$E2<>"")'], fill=red, stopIfTrue=True))
grey = PatternFill("solid", fgColor="D9D9D9")
s.conditional_formatting.add(f"E2:E{last}", FormulaRule(formula=['$D2="N"'], fill=grey, stopIfTrue=True))
green = PatternFill("solid", fgColor="C6EFCE")
s.conditional_formatting.add(f"D2:E{last}", FormulaRule(formula=['D2<>""'], fill=green))
s.freeze_panes = "D2"
s.auto_filter.ref = f"A1:H{last}"
for col, w in zip("ABCDEFGH", (8, 13, 30, 16, 26, 34, 10, 10)):
    s.column_dimensions[col].width = w

# ---------------- Progress ----------------
p = wb.create_sheet("Progress")
R = f"Rating!$D$2:$D${last}"
E = f"Rating!$E$2:$E${last}"
rows = [
    ("Items to rate", f"=COUNTA(Rating!$C$2:$C${last})"),
    ("Column D filled", f'=COUNTIF({R},"Y")+COUNTIF({R},"N")'),
    ("Column D answered Y", f'=COUNTIF({R},"Y")'),
    ("Column E filled (where D = Y)", f'=COUNTIFS({R},"Y",{E},"Y")+COUNTIFS({R},"Y",{E},"N")'),
    ("Rows finished", "=(B3-B4)+B5"),
    ("Percent complete", "=IF(B2=0,0,B6/B2)"),
    ("", None),
    ("Problems to fix (should all be 0)", None),
    ("D empty", "=B2-B3"),
    ("D = Y but E empty", "=B4-B5"),
    ("D = N but E filled (clear E)", f'=COUNTIFS({R},"N",{E},"<>")'),
    ("Scenes fully done (of 50)", f"=SUMPRODUCT(--(COUNTIFS(Rating!$A$2:$A${last},ROW($1:$50),Rating!$D$2:$D${last},\"\")=0))"),
]
p.cell(row=1, column=1, value="Progress (updates as you rate)").font = Font(name=F, size=14, bold=True)
for i, (k, f) in enumerate(rows, 2):
    p.cell(row=i, column=1, value=k).font = Font(name=F, bold=(f is None and k != ""))
    if f:
        c = p.cell(row=i, column=2, value=f)
        c.font = Font(name=F)
p["B7"].number_format = "0%"
p.column_dimensions["A"].width = 38
p.column_dimensions["B"].width = 14
red_font = PatternFill("solid", fgColor="F4B6B6")
p.conditional_formatting.add("B10:B12", FormulaRule(formula=["B10>0"], fill=red_font))
wb.active = 0
wb.save(out)
print("wrote", out, n, "rows,", len(scene_no), "scenes")
